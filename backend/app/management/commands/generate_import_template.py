# import openpyxl
# from django.core.management.base import BaseCommand

# class Command(BaseCommand):
#     def handle(self, *args, **kwargs):
#         wb = openpyxl.Workbook() #creates a new blank spreadsheet
#         ws = wb.active     # first sheet in the workbook
#         ws.title = "Products"  #rename sheet tab

#         #column headers - must match names the import view reads (row["name"]), row["price"],etc)
#         headers = ["name", "category", "price", "cost_price", "stock_count", "reorder_point"]
#         ws.append(headers)  # write header row

#        # example row so users see the expected format
#         ws.append(["Coca Cola 500ml", "Beverages", 60, 45, 48, 10])

#         wb.save("templates_data/product_import_template.xlsx")  # save file to disk
import openpyxl
from openpyxl.styles import Font
from django.core.management.base import BaseCommand

class Command(BaseCommand):
    def handle(self, *args, **kwargs):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products"

        headers = ["name", "category", "price", "cost_price", "stock_count", "reorder_point"]
        small_font = Font(size=10)

        ws.append(headers)

        # Apply smaller font to header row
        for cell in ws[1]:
            cell.font = small_font

        # Column widths
        widths = [22, 16, 10, 12, 14, 14]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

        wb.save("templates_data/product_import_template.xlsx")