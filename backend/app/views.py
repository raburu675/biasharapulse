import io
import openpyxl
from openpyxl.styles import Font
import pandas as pd
from django.http import FileResponse
from decimal import Decimal
from collections import defaultdict
from django.db.models import Sum, Count, Q, Max
from django.db.models.functions import TruncMonth, TruncWeek, TruncDay
from datetime import timedelta
from django.conf import settings
from django.utils import timezone
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.models import User

from rest_framework.decorators import (
    api_view, permission_classes, authentication_classes, throttle_classes,
)
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated, AllowAny
from rest_framework.authtoken.models import Token
from rest_framework.throttling import AnonRateThrottle

from .models import Business, Product, SaleRecord, Expense, StockMovement, Order, OrderItem
# PLAN: gating helpers (see plans.py)
from .plans import PLANS, require_feature, history_cutoff, sale_cap_response, sales_used
# AUTH: expiring tokens (see authentication.py)
from .authentication import fresh_token


TRUNC_FUNCS = {
    'daily': TruncDay,
    'weekly': TruncWeek,
    'monthly': TruncMonth,
}


# ═══════════════════════════════════════════════════════════
#  BUSINESS PLAN
# ═══════════════════════════════════════════════════════════

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def business_plan(request, business_id):
    """
    Tells the frontend what this business's plan allows, so it can
    show/lock pages and display usage.
    """
    try:
        business = Business.objects.get(id=business_id, owner=request.user)
    except Business.DoesNotExist:
        return Response({'error': 'Business not found'}, status=404)

    plan = business.plan
    return Response({
        'plan': business.plan_key,
        'subscription_tier': business.subscription_tier,
        'paid_until': business.paid_until,
        'trial_ends': business.trial_ends,
        'features': sorted(plan['features']),
        'limits': {
            'history_days': plan['history_days'],
            'max_users': plan['max_users'],
            'max_branches': plan['max_branches'],
            'total_sales_cap': plan['total_sales_cap'],
        },
        'usage': {
            'sales_used': sales_used(business),
        },
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def select_plan(request, business_id):
    """
    DEV ONLY: switches a business to the chosen tier without payment.
    Disabled when DEBUG is False. Replace with the M-Pesa success callback
    once payments exist.
    """
    if not settings.DEBUG:
        return Response({'error': 'Plan switching is disabled'}, status=403)

    try:
        business = Business.objects.get(id=business_id, owner=request.user)
    except Business.DoesNotExist:
        return Response({'error': 'Business not found'}, status=404)

    # Active paid plan → can't switch until it expires
    if business.plan_key not in ('starter', 'expired'):
        return Response({
            'error': 'You already have a paid plan. You can pick a new plan once your current plan expires.',
            'code': 'plan_active',
            'plan': business.plan_key,
            'paid_until': business.paid_until,
        }, status=409)

    tier = request.data.get('tier')
    if tier not in PLANS:
        return Response({'error': f'tier must be one of {list(PLANS)}'}, status=400)

    # Starter trial can only be selected once. After it expires, users
    # must pick a paid tier.
    if tier == 'starter' and business.trial_ends and business.trial_ends <= timezone.now():
        return Response({'error': 'Starter trial already used. Choose a paid plan.'}, status=400)

    business.subscription_tier = tier
    business.paid_until = None if tier == 'starter' else timezone.now() + timedelta(days=30)
    business.save(update_fields=['subscription_tier', 'paid_until'])

    return Response({
        'plan': business.plan_key,
        'paid_until': business.paid_until,
        'trial_ends': business.trial_ends,
    })


# ═══════════════════════════════════════════════════════════
#  DASHBOARD
# ═══════════════════════════════════════════════════════════

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_summary(request, business_id):
    """
    Executive dashboard data for one business.
    Accepts ?period=daily|weekly|monthly (defaults to monthly).
    """
    try:
        business = Business.objects.get(id=business_id, owner=request.user)
    except Business.DoesNotExist:
        return Response({'error': 'Business not found'}, status=404)

    period = request.GET.get('period', 'monthly')
    trunc_func = TRUNC_FUNCS.get(period, TruncMonth)

    sales = SaleRecord.objects.filter(business=business)
    expenses = Expense.objects.filter(business=business)
    products = Product.objects.filter(business=business)

    # PLAN: Starter only sees the last 30 days; expired sees no history
    cutoff = history_cutoff(business)
    if cutoff:
        sales = sales.filter(created_at__gte=cutoff)
        expenses = expenses.filter(created_at__gte=cutoff)

    net_revenue = sales.aggregate(total=Sum('amount'))['total'] or Decimal('0')
    total_expenses = expenses.aggregate(total=Sum('amount'))['total'] or Decimal('0')
    net_profit = max(Decimal('0'), net_revenue - total_expenses)
    net_margin = (net_profit / net_revenue * 100) if net_revenue > 0 else Decimal('0')
    active_inventory = products.aggregate(total=Sum('stock_count'))['total'] or 0

    payment_split = []
    if net_revenue > 0:
        by_channel = sales.values('payment_channel').annotate(total=Sum('amount'))
        for row in by_channel:
            percent = (row['total'] / net_revenue) * 100
            payment_split.append({
                'channel': row['payment_channel'],
                'percent': round(percent, 1),
            })

    category_volume = []
    total_stock = active_inventory
    if total_stock > 0:
        by_category = products.values('category').annotate(total=Sum('stock_count'))
        for row in by_category:
            percent = (row['total'] / total_stock) * 100
            category_volume.append({
                'category': row['category'],
                'percent': round(percent, 1),
            })

    period_sales = (
        sales.annotate(period=trunc_func('created_at'))
        .values('period')
        .annotate(total=Sum('amount'))
        .order_by('period')
    )
    period_expenses = (
        expenses.annotate(period=trunc_func('created_at'))
        .values('period')
        .annotate(total=Sum('amount'))
        .order_by('period')
    )

    period_data = defaultdict(lambda: {'sales': Decimal('0'), 'expenses': Decimal('0')})
    for row in period_sales:
        period_data[row['period']]['sales'] = row['total']
    for row in period_expenses:
        period_data[row['period']]['expenses'] = row['total']

    period_margin = []
    for p, data in sorted(period_data.items()):
        p_sales = data['sales']
        p_expenses = data['expenses']
        margin = ((p_sales - p_expenses) / p_sales * 100) if p_sales > 0 else Decimal('0')
        period_margin.append({
            'period': p,
            'margin': round(margin, 1),
        })

    recent_sales = sales.order_by('-created_at')[:5]
    recent_activity = [
        {
            'product': s.product.name if s.product else 'Unknown',
            'amount': s.amount,
            'payment_channel': s.payment_channel,
            'created_at': s.created_at,
        }
        for s in recent_sales
    ]

    return Response({
        'net_revenue': net_revenue,
        'expenses': total_expenses,
        'net_profit': net_profit,
        'net_margin': round(net_margin, 1),
        'active_inventory': active_inventory,
        'payment_channel_split': payment_split,
        'category_volume': category_volume,
        'period': period,
        'period_sales': [{'period': r['period'], 'total': r['total']} for r in period_sales],
        'period_expenses': [{'period': r['period'], 'total': r['total']} for r in period_expenses],
        'period_margin': period_margin,
        'recent_activity': recent_activity,
    })


# ═══════════════════════════════════════════════════════════
#  POS / PRODUCT PERFORMANCE
# ═══════════════════════════════════════════════════════════

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def pos_summary(request, business_id):
    """
    Per-product analytics: cost, price, stock, units sold, profit margin,
    sell-through rate, and performance tier for every product.
    """
    try:
        business = Business.objects.get(id=business_id, owner=request.user)
    except Business.DoesNotExist:
        return Response({'error': 'Business not found'}, status=404)

    # PLAN GATE: pos feature
    denied = require_feature(business, 'pos')
    if denied:
        return denied

    # PLAN: Product performance analytics start at Growth.
    can_scan = business.has_feature('qr_scanning')
    if not business.has_feature('product_performance'):
        lean = []
        for p in Product.objects.filter(business=business):
            if p.stock_count <= 0:
                stock_status = 'Out of Stock'
            elif p.stock_count <= p.reorder_point:
                stock_status = 'Low Stock'
            else:
                stock_status = 'In Stock'
            lean.append({
                'id': p.id,
                'name': p.name,
                'category': p.category,
                'barcode': p.barcode if can_scan else None,
                'selling_price': p.price,
                'stock_quantity': p.stock_count,
                'stock_status': stock_status,
            })
        return Response({'products': lean, 'analytics_locked': True})

    products = Product.objects.filter(business=business)

    product_list = []
    for p in products:
        sales = SaleRecord.objects.filter(business=business, product=p)
        units_sold = sales.aggregate(total=Sum('quantity'))['total'] or 0
        total_revenue = sales.aggregate(total=Sum('amount'))['total'] or Decimal('0')

        total_cost = Decimal(units_sold) * p.cost_price
        gross_profit = total_revenue - total_cost
        profit_margin = (gross_profit / total_revenue * 100) if total_revenue > 0 else Decimal('0')

        total_handled = units_sold + p.stock_count
        sell_through_rate = (Decimal(units_sold) / Decimal(total_handled) * 100) if total_handled > 0 else Decimal('0')

        if sell_through_rate >= 70:
            performance_tier = 'Star Performer'
        elif sell_through_rate >= 40:
            performance_tier = 'Steady'
        else:
            performance_tier = 'Slow Mover'

        if p.stock_count <= 0:
            stock_status = 'Out of Stock'
        elif p.stock_count <= p.reorder_point:
            stock_status = 'Low Stock'
        else:
            stock_status = 'In Stock'

        product_list.append({
            'id': p.id,
            'name': p.name,
            'category': p.category,
            'cost_price': p.cost_price,
            'barcode': p.barcode if can_scan else None,
            'selling_price': p.price,
            'stock_quantity': p.stock_count,
            'units_sold': units_sold,
            'total_revenue': total_revenue,
            'total_cost': total_cost,
            'gross_profit': gross_profit,
            'profit_margin': round(profit_margin, 1),
            'sell_through_rate': round(sell_through_rate, 1),
            'performance_tier': performance_tier,
            'stock_status': stock_status,
        })

    total_inventory_value = sum((item['stock_quantity'] * item['cost_price'] for item in product_list), Decimal('0'))
    total_revenue = sum((item['total_revenue'] for item in product_list), Decimal('0'))
    total_profit = sum((item['gross_profit'] for item in product_list), Decimal('0'))
    avg_sell_through = (
        sum((item['sell_through_rate'] for item in product_list), Decimal('0')) / len(product_list)
        if product_list else Decimal('0')
    )

    top_seller = max(product_list, key=lambda i: i['units_sold']) if product_list else None
    most_profitable = max(product_list, key=lambda i: i['gross_profit']) if product_list else None
    best_margin = max(product_list, key=lambda i: i['profit_margin']) if product_list else None

    return Response({
        'products': product_list,
        'total_inventory_value': total_inventory_value,
        'total_revenue': total_revenue,
        'total_profit': total_profit,
        'avg_sell_through': round(avg_sell_through, 1),
        'top_seller': top_seller,
        'most_profitable': most_profitable,
        'best_margin': best_margin,
    })

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def product_performance(request, business_id):
    """
    Per-product performance over a period (?days=7|30|90): velocity,
    days of stock cover, last sale, margin and all-time totals.
    """
    try:
        business = Business.objects.get(id=business_id, owner=request.user)
    except Business.DoesNotExist:
        return Response({'error': 'Business not found'}, status=404)

    # PLAN GATE: Growth and above
    denied = require_feature(business, 'product_performance')
    if denied:
        return denied

    try:
        days = int(request.GET.get('days', 30))
    except ValueError:
        days = 30
    days = min(max(days, 1), 365)

    now = timezone.now()
    since = now - timedelta(days=days)
    sales = SaleRecord.objects.filter(business=business)

    # All-time totals and last sale date, per product
    totals = {
        r['product']: r
        for r in sales.values('product').annotate(
            units=Sum('quantity'), revenue=Sum('amount'), last=Max('created_at')
        )
    }
    # Totals inside the selected period
    recent = {
        r['product']: r
        for r in sales.filter(created_at__gte=since).values('product').annotate(
            units=Sum('quantity'), revenue=Sum('amount')
        )
    }

    data = []
    for p in Product.objects.filter(business=business):
        t = totals.get(p.id, {})
        r = recent.get(p.id, {})

        units_total = t.get('units') or 0
        revenue_total = t.get('revenue') or Decimal('0')
        units_period = r.get('units') or 0
        revenue_period = r.get('revenue') or Decimal('0')
        last_sale = t.get('last')

        price = Decimal(str(p.price))
        gross_profit = revenue_total - Decimal(units_total) * p.cost_price

        # Realised margin once there are sales, otherwise the unit margin
        if revenue_total > 0:
            margin = gross_profit / revenue_total * 100
        elif price > 0:
            margin = (price - p.cost_price) / price * 100
        else:
            margin = Decimal('0')

        handled = units_total + p.stock_count
        sell_through = (Decimal(units_total) / Decimal(handled) * 100) if handled > 0 else Decimal('0')

        velocity = units_period / days                      # units per day
        cover = round(p.stock_count / velocity, 1) if velocity > 0 else None

        if p.stock_count <= 0:
            stock_status = 'Out of Stock'
        elif p.stock_count <= p.reorder_point:
            stock_status = 'Low Stock'
        else:
            stock_status = 'In Stock'

        data.append({
            'id': p.id,
            'name': p.name,
            'category': p.category,
            'cost_price': p.cost_price,
            'selling_price': p.price,
            'stock_quantity': p.stock_count,
            'reorder_point': p.reorder_point,
            'stock_status': stock_status,
            'units_period': units_period,
            'revenue_period': revenue_period,
            'units_total': units_total,
            'revenue_total': revenue_total,
            'gross_profit': gross_profit,
            'profit_margin': round(margin, 1),
            'sell_through_rate': round(sell_through, 1),
            'velocity': round(velocity, 2),
            'days_of_cover': cover,
            'last_sale_at': last_sale,
            'days_since_last_sale': (now - last_sale).days if last_sale else None,
        })

    return Response({'days': days, 'products': data})


# ═══════════════════════════════════════════════════════════
#  STOCK MOVEMENTS
# ═══════════════════════════════════════════════════════════

@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated])
def stock_movements(request, business_id):
    """
    GET  -> return the recent stock movement log
    POST -> create a new Stock In / Waste-Damage adjustment
    """
    try:
        business = Business.objects.get(id=business_id, owner=request.user)
    except Business.DoesNotExist:
        return Response({'error': 'Business not found'}, status=404)

    # PLAN GATE: stock_movement feature
    denied = require_feature(business, 'stock_movement')
    if denied:
        return denied

    if request.method == 'GET':
        movement_type = request.GET.get('type')

        movements = (
            StockMovement.objects.filter(business=business)
            .select_related('product', 'user')
            .order_by('-created_at')
        )

        cutoff = history_cutoff(business)
        if cutoff:
            movements = movements.filter(created_at__gte=cutoff)

        if movement_type:
            movements = movements.filter(movement_type=movement_type)

        movements = movements[:20]

        log_data = []
        for m in movements:
            if m.movement_type == 'stock_in':
                qty_display = f"+{m.quantity_change} units"
            elif m.movement_type == 'waste_damage':
                qty_display = f"{m.quantity_change} units ({m.resulting_stock} left)"
            else:
                qty_display = f"{m.resulting_stock} units left"

            log_data.append({
                'id': m.id,
                'type': m.get_movement_type_display(),
                'item': m.product.name if m.product else 'Unknown',
                'qty': qty_display,
                'time': m.created_at,
                'user': m.user.username if m.user else 'System',
                'category': m.product.category if m.product else '',
                'currentStock': m.resulting_stock,
                'reorderPoint': m.product.reorder_point if m.product else 0,
            })

        return Response({'movements': log_data})

    product_id = request.data.get('product_id')
    movement_type = request.data.get('movement_type')
    quantity_change = request.data.get('quantity_change')
    note = request.data.get('note', '')

    if not product_id or not movement_type or quantity_change is None:
        return Response({'error': 'product_id, movement_type, and quantity_change are required'}, status=400)

    if movement_type not in ('stock_in', 'waste_damage'):
        return Response({'error': "movement_type must be 'stock_in' or 'waste_damage'"}, status=400)

    try:
        product = Product.objects.get(id=product_id, business=business)
    except Product.DoesNotExist:
        return Response({'error': 'Product not found'}, status=404)

    quantity_change = int(quantity_change)
    if movement_type == 'waste_damage' and quantity_change > 0:
        quantity_change = -quantity_change

    product.stock_count = max(0, product.stock_count + quantity_change)
    product.save(update_fields=['stock_count'])

    movement = StockMovement.objects.create(
        business=business,
        product=product,
        movement_type=movement_type,
        quantity_change=quantity_change,
        resulting_stock=product.stock_count,
        note=note,
    )

    # Low-stock-alert check for manual adjustments only lives here.
    if product.stock_count <= product.reorder_point:
        StockMovement.objects.create(
            business=business,
            product=product,
            movement_type='low_stock_alert',
            quantity_change=0,
            resulting_stock=product.stock_count,
            note=f'Stock at {product.stock_count}, reorder point is {product.reorder_point}',
        )

    return Response({
        'id': movement.id,
        'product': product.name,
        'movement_type': movement.get_movement_type_display(),
        'quantity_change': movement.quantity_change,
        'resulting_stock': movement.resulting_stock,
        'created_at': movement.created_at,
    }, status=201)


# ═══════════════════════════════════════════════════════════
#  SALES
# ═══════════════════════════════════════════════════════════

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def create_sale(request, business_id):
    """
    Records a sale. Stock decrement, amount calculation, and the low-stock
    alert all happen inside SaleRecord.save().
    """
    try:
        business = Business.objects.get(id=business_id, owner=request.user)
    except Business.DoesNotExist:
        return Response({'error': 'Business not found'}, status=404)

    # PLAN GATE: pos feature (expired plans can't record sales)
    denied = require_feature(business, 'pos')
    if denied:
        return denied

    # PLAN LIMIT: Starter is limited to its first 100 sales
    capped = sale_cap_response(business)
    if capped:
        return capped

    product_id = request.data.get('product_id')
    quantity = request.data.get('quantity', 1)
    payment_channel = request.data.get('payment_channel')

    if not product_id or not payment_channel:
        return Response({'error': 'product_id and payment_channel are required'}, status=400)

    try:
        product = Product.objects.get(id=product_id, business=business)
    except Product.DoesNotExist:
        return Response({'error': 'Product not found'}, status=404)

    quantity = int(quantity)
    if quantity > product.stock_count:
        return Response({'error': f'Only {product.stock_count} units in stock'}, status=400)

    sale = SaleRecord.objects.create(
        business=business,
        product=product,
        quantity=quantity,
        payment_channel=payment_channel,
    )

    product.refresh_from_db()

    return Response({
        'id': sale.id,
        'product': product.name,
        'amount': sale.amount,
        'quantity': sale.quantity,
        'payment_channel': sale.payment_channel,
        'remaining_stock': product.stock_count,
        'created_at': sale.created_at,
    }, status=201)


# ═══════════════════════════════════════════════════════════
#  ORDERS
# ═══════════════════════════════════════════════════════════

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def order_list(request, business_id):
    """
    Returns all orders for a business, newest first, with line items and total.
    """
    try:
        business = Business.objects.get(id=business_id, owner=request.user)
    except Business.DoesNotExist:
        return Response({'error': 'Business not found'}, status=404)

    # PLAN GATE: orders feature
    denied = require_feature(business, 'orders')
    if denied:
        return denied

    orders = Order.objects.filter(business=business).prefetch_related('items').order_by('-created_at')

    order_data = []
    for o in orders:
        order_data.append({
            'id': o.id,
            'order_number': o.order_number,
            'customer_name': o.customer_name,
            'customer_phone': o.customer_phone,
            'shipping_address': o.shipping_address,
            'items': [
                {'name': item.name, 'qty': item.quantity}
                for item in o.items.all()
            ],
            'total_amount': o.total_amount,
            'payment_method': o.get_payment_method_display(),
            'created_at': o.created_at,
            'source': o.get_source_display(),
            'courier': o.courier,
            'tracking_number': o.tracking_number,
            'status': o.status,
        })

    return Response({'orders': order_data})


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def update_order_status(request, business_id, order_id):
    """
    Updates an order's status. Advancing to 'delivered' triggers
    Order.save() to create SaleRecords for each line item.
    """
    try:
        business = Business.objects.get(id=business_id, owner=request.user)
    except Business.DoesNotExist:
        return Response({'error': 'Business not found'}, status=404)

    # PLAN GATE: orders feature
    denied = require_feature(business, 'orders')
    if denied:
        return denied

    try:
        order = Order.objects.get(id=order_id, business=business)
    except Order.DoesNotExist:
        return Response({'error': 'Order not found'}, status=404)

    new_status = request.data.get('status')
    valid_statuses = [choice[0] for choice in Order.STATUS_CHOICES]
    if new_status not in valid_statuses:
        return Response({'error': f'status must be one of {valid_statuses}'}, status=400)

    # PLAN LIMIT: delivering counts toward the Starter sale allowance
    if new_status == 'delivered' and order.status != 'delivered':
        capped = sale_cap_response(business)
        if capped:
            return capped

    order.status = new_status
    order.save()

    return Response({
        'id': order.id,
        'order_number': order.order_number,
        'status': order.status,
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def create_order(request, business_id):
    """
    Creates a customer order with real Product-backed line items.
    """
    try:
        business = Business.objects.get(id=business_id, owner=request.user)
    except Business.DoesNotExist:
        return Response({'error': 'Business not found'}, status=404)

    # PLAN GATE: orders feature
    denied = require_feature(business, 'orders')
    if denied:
        return denied

    customer_name = request.data.get('customer_name')
    items = request.data.get('items', [])

    if not customer_name or not items:
        return Response({'error': 'customer_name and items are required'}, status=400)

    order = Order.objects.create(
        business=business,
        customer_name=customer_name,
        customer_phone=request.data.get('customer_phone', ''),
        shipping_address=request.data.get('shipping_address', ''),
        payment_method=request.data.get('payment_method', 'mpesa'),
        source=request.data.get('source', 'website'),
        courier=request.data.get('courier', ''),
        tracking_number=request.data.get('tracking_number', ''),
        status='pending',
    )

    for item in items:
        try:
            product = Product.objects.get(id=item['product_id'], business=business)
        except Product.DoesNotExist:
            continue
        OrderItem.objects.create(
            order=order,
            product=product,
            quantity=item.get('quantity', 1),
        )

    return Response({
        'id': order.id,
        'order_number': order.order_number,
        'total_amount': order.total_amount,
    }, status=201)


# ═══════════════════════════════════════════════════════════
#  PRODUCT IMPORT
# ═══════════════════════════════════════════════════════════

COLUMN_ALIASES = {
    'product name': 'name',
    'selling price (kes)': 'price',
    'cost price (kes)': 'cost_price',
    'stock count': 'stock_count',
    'reorder point': 'reorder_point',
}


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def import_products(request, business_id):
    """
    Bulk create/update Products from an uploaded spreadsheet (.xlsx or .csv).
    """
    try:
        business = Business.objects.get(id=business_id, owner=request.user)
    except Business.DoesNotExist:
        return Response({'error': 'Business not found'}, status=404)

    file = request.FILES.get('file')
    if not file:
        return Response({'error': 'No file uploaded'}, status=400)

    try:
        if file.name.endswith('.csv'):
            df = pd.read_csv(file)
        else:
            df = pd.read_excel(file)
    except Exception as e:
        return Response({'error': f'Could not read file: {e}'}, status=400)

    df.columns = [str(c).strip().lower() for c in df.columns]
    df = df.rename(columns=COLUMN_ALIASES)

    required = {'name', 'category', 'price'}
    missing = required - set(df.columns)
    if missing:
        return Response({
            'error': f"Missing columns: {', '.join(missing)}. Download the template and use those exact column names.",
            'found_columns': list(df.columns),
        }, status=400)

    created, updated, errors = 0, 0, []

    for i, row in df.iterrows():
        try:
            name = str(row['name']).strip()

            product, was_created = Product.objects.update_or_create(
                business=business,
                name__iexact=name,
                defaults={
                    'name': name,
                    'category': str(row['category']).strip(),
                    'price': float(row['price']),
                    'cost_price': float(row.get('cost_price', 0) or 0),
                    'stock_count': int(row.get('stock_count', 0) or 0),
                    'reorder_point': int(row.get('reorder_point', 5) or 5),
                },
            )
            created += was_created
            updated += not was_created
        except Exception as e:
            errors.append({'row': int(i) + 2, 'error': str(e)})

    return Response({'created': created, 'updated': updated, 'errors': errors})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def import_template(request):
    """
    Generates the blank product import template on the fly and streams it
    straight back.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"

    headers = ["Product Name", "Category", "Selling Price (KES)", "Cost Price (KES)", "Stock Count", "Reorder Point"]
    small_font = Font(size=10)

    ws.append(headers)
    for cell in ws[1]:
        cell.font = small_font

    widths = [24, 16, 18, 16, 14, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return FileResponse(buffer, as_attachment=True, filename='product_import_template.xlsx')


# ═══════════════════════════════════════════════════════════
#  AUTH
# ═══════════════════════════════════════════════════════════

class SignupThrottle(AnonRateThrottle):
    """Limits sign-ups per IP so one person can't farm free trials."""
    scope = 'signup'

    def get_rate(self):
        return '100/hour' if settings.DEBUG else '5/hour'


def canonical_email(email):
    """
    Collapses tricks that point to the same inbox:
    a.b+x@gmail.com -> ab@gmail.com, name+x@other.com -> name@other.com
    """
    local, _, domain = email.lower().strip().partition('@')
    local = local.split('+')[0]
    if domain in ('gmail.com', 'googlemail.com'):
        local = local.replace('.', '')
        domain = 'gmail.com'
    return f'{local}@{domain}'


def email_taken(email):
    # DEBUG: exact match only, so you can test with a+1@gmail.com, a+2@gmail.com
    if settings.DEBUG:
        return User.objects.filter(email__iexact=email).exists()

    target = canonical_email(email)
    domain = target.split('@')[1]
    domain_q = Q(email__iendswith='@gmail.com') | Q(email__iendswith='@googlemail.com') \
        if domain == 'gmail.com' else Q(email__iendswith='@' + domain)

    return any(
        canonical_email(e) == target
        for e in User.objects.filter(domain_q).values_list('email', flat=True)
    )


@api_view(['POST'])
@authentication_classes([])   # public: a stale token must not block signup
@permission_classes([AllowAny])
@throttle_classes([SignupThrottle])
def signup(request):
    """
    Creates a User and their first Business in one shot, then returns a token.
    Public — no auth required.
    """
    username = (request.data.get('username') or '').strip()
    email = (request.data.get('email') or '').strip().lower()
    password = request.data.get('password') or ''
    business_name = (request.data.get('business_name') or '').strip()

    if not username or not password:
        return Response({'error': 'username and password are required'}, status=400)

    if not email:
        return Response({'error': 'Email is required'}, status=400)

    try:
        validate_email(email)
    except ValidationError:
        return Response({'error': 'Enter a valid email address'}, status=400)

    if len(password) < 8:
        return Response({'error': 'Password must be at least 8 characters'}, status=400)

    if User.objects.filter(username__iexact=username).exists():
        return Response({'error': 'That username is already taken'}, status=400)

    if email_taken(email):
        return Response({'error': 'That email is already registered'}, status=400)

    user = User.objects.create_user(username=username, email=email, password=password)

    # Every user gets a business on day one, with a 30-day Starter trial
    business = Business.objects.create(
        owner=user,
        name=business_name or f"{username}'s Business",
        trial_ends=timezone.now() + timedelta(days=30),
    )

    token = fresh_token(user)
    login(request, user)

    return Response({
        'token': token.key,
        'businessId': business.id,
        'username': user.username,
        'business_name': business.name,
    }, status=201)


@api_view(['POST'])
@authentication_classes([])   # public: a stale token must not block login
@permission_classes([AllowAny])
def login_view(request):
    """
    Verifies credentials, returns the same payload shape as signup.
    Public — no auth required.
    """
    username = (request.data.get('username') or '').strip()
    password = request.data.get('password') or ''

    user = authenticate(request, username=username, password=password)
    if user is None:
        return Response({'error': 'Invalid username or password'}, status=401)

    login(request, user)
    token = fresh_token(user)

    business = user.businesses.first()
    if business is None:
        business = Business.objects.create(
            owner=user,
            name=f"{user.username}'s Business",
            trial_ends=timezone.now() + timedelta(days=30),
        )

    return Response({
        'token': token.key,
        'businessId': business.id,
        'username': user.username,
        'business_name': business.name,
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def logout_view(request):
    """
    Deletes the current token so it can't be reused, and clears the session.
    """
    Token.objects.filter(user=request.user).delete()
    logout(request)
    return Response({'ok': True})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def me(request):
    """
    Returns the current user's info + their businesses.
    """
    businesses = list(request.user.businesses.values('id', 'name', 'subscription_tier'))
    return Response({
        'username': request.user.username,
        'email': request.user.email,
        'businesses': businesses,
    })